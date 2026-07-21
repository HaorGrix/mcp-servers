"""Render a PRD content dict into a .docx matching the HaorGrix PRD template.

The template is used as the style donor: we open it, strip its body, and rebuild
the document so headings, table styles, fonts and theme carry over unchanged.
"""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

GUIDANCE_GREY = RGBColor(0x80, 0x80, 0x80)


class PRDBuildError(RuntimeError):
    """Raised when the template is missing a style the renderer depends on."""


def _clear_body(doc: Document) -> None:
    body = doc.element.body
    sect = copy.deepcopy(body.xpath("./w:sectPr")[0]) if body.xpath("./w:sectPr") else None
    for child in list(body):
        body.remove(child)
    if sect is not None:
        body.append(sect)


def _style_map(doc: Document) -> dict:
    """Map display name -> style object.

    python-docx's `styles[name]` lookup does not reliably resolve built-in
    heading styles in this template, so resolve against the actual style
    collection instead.
    """
    return {s.name: s for s in doc.styles}


def _style_or_fail(styles: dict, name: str):
    style = styles.get(name)
    if style is None:
        raise PRDBuildError(f"template is missing required style {name!r}")
    return style


def _table_style(styles: dict):
    for candidate in ("Table Grid", "TableGrid", "Light Grid Accent 1"):
        if candidate in styles:
            return styles[candidate]
    raise PRDBuildError("template has no usable table style")


class PRDWriter:
    def __init__(self, template: Path):
        self.doc = Document(str(template))
        _clear_body(self.doc)
        styles = _style_map(self.doc)
        self.table_style = _table_style(styles)
        self.styles = {
            key: _style_or_fail(styles, key)
            for key in ("Title", "Heading 1", "Heading 2", "List Bullet")
        }

    # --- primitives -----------------------------------------------------
    def title(self, text: str, subtitle: str = "", tagline: str = "") -> None:
        p = self.doc.add_paragraph(text, style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for line, size, bold in ((subtitle, 16, True), (tagline, 11, False)):
            if not line:
                continue
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = q.add_run(line)
            run.bold = bold
            run.font.size = Pt(size)

    def h1(self, text: str) -> None:
        self.doc.add_paragraph(text, style="Heading 1")

    def h2(self, text: str) -> None:
        self.doc.add_paragraph(text, style="Heading 2")

    def para(self, text: str) -> None:
        self.doc.add_paragraph(text)

    def meta(self, text: str) -> None:
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.italic = True
        run.font.color.rgb = GUIDANCE_GREY
        run.font.size = Pt(9)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def table(self, header: list[str], rows: list[list[str]]) -> None:
        if not header:
            raise PRDBuildError("table requires a header row")
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = self.table_style
        for cell, label in zip(t.rows[0].cells, header):
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
        for row in rows:
            if len(row) != len(header):
                raise PRDBuildError(f"row {row!r} does not match header width {len(header)}")
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
        self.doc.add_paragraph()

    def page_break(self) -> None:
        self.doc.add_page_break()

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


def render(content: dict, template: Path, out_path: Path) -> Path:
    """Render a PRD content dict.

    `content` keys: title, subtitle, tagline, then `sections` — a list of
    blocks, each a dict with a `kind` of h1/h2/para/meta/bullets/table/break.
    """
    w = PRDWriter(template)
    w.title(content["title"], content.get("subtitle", ""), content.get("tagline", ""))

    handlers = {
        "h1": lambda b: w.h1(b["text"]),
        "h2": lambda b: w.h2(b["text"]),
        "para": lambda b: w.para(b["text"]),
        "meta": lambda b: w.meta(b["text"]),
        "bullets": lambda b: w.bullets(b["items"]),
        "table": lambda b: w.table(b["header"], b["rows"]),
        "break": lambda b: w.page_break(),
    }
    for block in content["sections"]:
        kind = block.get("kind")
        handler = handlers.get(kind)
        if handler is None:
            raise PRDBuildError(f"unknown block kind {kind!r}")
        handler(block)

    return w.save(out_path)
